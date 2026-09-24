import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

d = Document('report_corrected.docx')

# ---------- helpers ----------
def replace_text(p, old, new):
    runs = p.runs
    full = ''.join(r.text for r in runs)
    i = full.find(old)
    if i < 0: return False
    j = i + len(old)
    pos = 0; done = False
    for r in runs:
        a, b = pos, pos + len(r.text); pos = b
        if b <= i or a >= j: continue
        s = max(a, i) - a; e = min(b, j) - a
        if not done:
            r.text = r.text[:s] + new + r.text[e:]; done = True
        else:
            r.text = r.text[:s] + r.text[e:]
    return True

def plain(text):
    r = OxmlElement('w:r'); t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve'); t.text = text; r.append(t); return r

def add_links(p, urls):
    part = p.part
    for u in urls:
        rid = part.relate_to(u, RT.HYPERLINK, is_external=True)
        p._p.append(plain(' ')); p._p.append(plain('·')); p._p.append(plain(' '))
        hl = OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rid)
        r = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
        st = OxmlElement('w:rStyle'); st.set(qn('w:val'), 'Hyperlink'); rPr.append(st)
        un = OxmlElement('w:u'); un.set(qn('w:val'), 'single'); rPr.append(un)
        r.append(rPr)
        t = OxmlElement('w:t'); t.set(qn('xml:space'), 'preserve'); t.text = u
        r.append(t); hl.append(r); p._p.append(hl)

# ---------- link additions ----------
ADD = {
 18: ['https://stayathomemacro.substack.com/p/sahm-thing-more-on-the-sahm-rule',
      'https://www.bloomberg.com/opinion/articles/2024-08-07/the-sahm-rule-is-warning-of-recession-but-claudia-sahm-isn-t-sold'],
 62: ['https://cepr.org/voxeu/columns/making-sense-conflicting-labour-market-signals'],
 64: ['https://www.chicagofed.org/publications/chicago-fed-insights/2026/low-hire-low-fire-chicago-fed-labor-market-indicators'],
 68: ['https://www.federalreserve.gov/econres/notes/feds-notes/population-growth-and-labor-market-fragility-lessons-from-domestic-and-international-experiences-20260520.html'],
 70: ['https://www.bloomberg.com/opinion/articles/2024-08-07/the-sahm-rule-is-warning-of-recession-but-claudia-sahm-isn-t-sold'],
 78: ['https://stayathomemacro.substack.com/p/sahm-thing-more-on-the-sahm-rule',
      'https://budgetlab.yale.edu/research/what-we-learned-about-recession-indicators-rules-and-triggers-after-pandemic'],
 92: ['https://www.research.ed.ac.uk/files/8518774/ELSBY_2012_The_ins_and_outs_on_cyclical_unemployment.pdf'],
 128:['https://cdhowe.org/publication/dont-be-too-quick-to-call-a-recession/',
      'https://insurancenewsnet.com/innarticle/do-2-down-quarters-mean-recession-not-necessarily'],
 178:['https://eml.berkeley.edu/~cromer/Working/NBER%20Recession%20Dates/NBER%20Recession%20Dates.pdf'],
 228:['https://www.federalreserve.gov/econres/notes/feds-notes/dont-fear-the-yield-curve-reprise-20220325.html'],
 272:['https://www.richmondfed.org/-/media/RichmondFedOrg/publications/research/economic_quarterly/2003/summer/pdf/stockwatsonsummer03.pdf'],
 274:['https://www.princeton.edu/~mwatson/papers/Stock_Watson_Predicting_Recessions_1993.pdf'],
 288:['https://ideas.repec.org/a/tpr/restat/v80y1998i1p45-61.html'],
 314:['https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/coincident/state-business-cycle-dates.xlsx',
      'https://www.philadelphiafed.org/-/media/FRBP/Assets/Surveys-And-Data/coincident/state-business-cycle-update-highlights.pdf'],
 332:['https://conversableeconomist.com/2022/07/25/is-a-recession-defined-as-two-negative-quarters/'],
 348:['https://conversableeconomist.com/wp-content/uploads/2023/02/image-3.png'],
 352:['https://www.brookings.edu/wp-content/uploads/1991/06/1991b_bpea_bernanke_lown_friedman.pdf'],
}
for idx, urls in ADD.items():
    p = d.paragraphs[idx]
    assert p.text.strip().startswith('http'), (idx, p.text[:60])
    add_links(p, urls)

# ---------- text edits ----------
# 1. Blinder source note
ok = replace_text(d.paragraphs[347],
 'Table 1 read from the publisher’s typeset PDF, recovered from the Internet Archive capture of pubs.aeaweb.org (the live AEA PDF returns 403)',
 'Table 1 read from the publisher’s typeset table, reproduced as an image in Timothy Taylor’s Conversable Economist post of 8 February 2023 and linked below; the AEA PDF itself returns 403 to every automated route')
assert ok, 'blinder note'

# 2. Elsby-Michaels-Solon: record the working-paper divergence
ok = replace_text(d.paragraphs[91],
 'Both sentences are quoted from the published American Economic Journal: Macroeconomics text, read at the University of Edinburgh’s repository copy.',
 'Both sentences are quoted from the published American Economic Journal: Macroeconomics text, read at the University of Edinburgh’s repository copy, which is linked below and is the article’s only open-access location. The working-paper versions — NBER 12853 and the January 2007 Michigan draft — carry the second sentence without its parenthetical, as “Shimer’s claim that the inflow rate is ‘nearly acyclical’ is an overstatement at best”, so the NBER link below will not reproduce the published wording.')
assert ok, 'ems note'

# 3. Shiskin: retire the unverifiable 1.5-rendering clause
ok = replace_text(d.paragraphs[127],
 'the intended value is almost certainly 1.5, misprinted in the original, and a second, independent reproduction of the column does render the figure as 1.5 per cent.',
 'the intended value is almost certainly 1.5, misprinted in the original. Three independent reproductions of the column — Ritholtz (2020), the C.D. Howe Institute’s, and the syndicated newspaper version linked below — all render the figure as 15 per cent, and two of the three preserve the column’s own missing preposition, which places the error in the Times’s text rather than in any later transcription of it. No reproduction rendering the figure as 1.5 per cent was located.')
assert ok, 'shiskin note'

d.save('report_corrected.docx')
print('saved')
