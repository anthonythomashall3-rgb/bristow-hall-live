import copy, re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT

d=Document('report.docx')

def replace_across_runs(p, old, new):
    runs=p.runs
    full=''.join(r.text for r in runs)
    i=full.find(old)
    if i<0: return False
    # map char positions to runs
    pos=0; spans=[]
    for r in runs:
        spans.append((pos,pos+len(r.text),r)); pos+=len(r.text)
    j=i+len(old)
    for a,b,r in spans:
        if b<=i or a>=j: continue
        s=max(a,i)-a; e=min(b,j)-a
        head=r.text[:s]; tail=r.text[e:]
        r.text = head + (new if a<=i<b else '') + tail
        if not (a<=i<b): r.text = head + tail
    return True

# ---- D3: entry [41] author order
p=d.paragraphs[41]
assert replace_across_runs(p, 'Kurt G. Lunsford, Pawel M. Krolikowski and Meifeng Yang',
                              'Pawel M. Krolikowski, Kurt G. Lunsford and Meifeng Yang'), 'D3 failed'

# ---- D2: entry [91] drop false quotation marks
p=d.paragraphs[91]
ok = replace_across_runs(p, 'and that “the separation rate leads unemployment, while the job finding rate moves contemporaneously with it”',
                            'and that the separation rate leads unemployment while the job finding rate moves contemporaneously with it')
assert ok, 'D2 failed'

# ---- D1: entry [18] add the chapter URL
p=d.paragraphs[18]
part=p.part
rid=part.relate_to('https://www.hamiltonproject.org/assets/files/Sahm_web_20190506.pdf', RT.HYPERLINK, is_external=True)
def plain(text):
    r=OxmlElement('w:r'); t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve'); t.text=text; r.append(t); return r
hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rid)
r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
st=OxmlElement('w:rStyle'); st.set(qn('w:val'),'Hyperlink'); rPr.append(st)
u=OxmlElement('w:u'); u.set(qn('w:val'),'single'); rPr.append(u)
r.append(rPr)
t=OxmlElement('w:t'); t.set(qn('xml:space'),'preserve')
t.text='https://www.hamiltonproject.org/assets/files/Sahm_web_20190506.pdf'
r.append(t); hl.append(r)
# insert after the FIRST hyperlink (the summary), so order is summary, chapter, FRED
first=p._p.findall(qn('w:hyperlink'))[0]
first.addnext(hl)
hl.addprevious(plain(' ')) if False else None
# add " · " separators between first and new
sep_sp=plain(' '); sep_dot=plain('·'); sep_sp2=plain(' ')
first.addnext(sep_sp2); first.addnext(sep_dot); first.addnext(sep_sp)
d.save('report_corrected.docx')
print('saved report_corrected.docx')
