import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy
d = Document('report_corrected.docx')

def replace_text(p, old, new):
    runs=p.runs; full=''.join(r.text for r in runs); i=full.find(old)
    if i<0: return False
    j=i+len(old); pos=0; done=False
    for r in runs:
        a,b=pos,pos+len(r.text); pos=b
        if b<=i or a>=j: continue
        s=max(a,i)-a; e=min(b,j)-a
        if not done: r.text=r.text[:s]+new+r.text[e:]; done=True
        else: r.text=r.text[:s]+r.text[e:]
    return True

# --- D: shutdown start date. CRS: "began on October 1", concluded November 12, 2025.
assert replace_text(d.paragraphs[371],
  '42 full days from 30 September to 12 November 2025',
  '42 full days from 1 October to 12 November 2025'), 'shutdown date'

# --- add a dated entry to the corrections log recording this audit
anchor = d.paragraphs[458]     # "Note on how to read this log..."
new_p = copy.deepcopy(anchor._p)
anchor._p.addnext(new_p)
from docx.text.paragraph import Paragraph
np = Paragraph(new_p, anchor._parent)
for r in np.runs[1:]: r.text=''
np.runs[0].text = (
 "External audit of August 30, 2026, applied in the body above. An outside pass re-verified the report against "
 "sources fetched anew and applied thirty-six changes, none of which alters a substantive finding. Three defects: "
 "the Sahm-rule entry linked the two-page proposal summary rather than the twenty-six-page chapter it quotes three "
 "times (the chapter is now linked, and all three quotations verify against it); the labour-flow entry printed a "
 "spliced paraphrase of Fujita and Ramey as a quotation, and the quotation marks are removed; and the WARN-Act "
 "entry reversed the first two authors of Cleveland Fed Economic Commentary 2019-21, whose published order is "
 "Krolikowski, Lunsford and Yang. Three sentences carried a period both inside and after the closing quotation "
 "mark, two of them attaching a terminal period to a title that has none. The 2025 funding gap ran from 1 October, "
 "not 30 September, on the register's own wording. One unverifiable specific was retired under the standing rule: "
 "the claim that a second reproduction of Shiskin's column renders the employment figure as 1.5 per cent, which no "
 "route could confirm and which four readable reproductions contradict; the misprint argument now rests on the "
 "arithmetic, independently recomputed from FRED's PAYEMS as 1.46 to 5.19 per cent across the six recessions from "
 "1948. The Camacho, Perez-Quiros and Poncela citation gained its volume, issue and pages. The remaining changes "
 "are twenty-five source links added to entries that quoted a document they did not link. Two live-table cells "
 "carry a later revision in parentheses: initial claims and the CFNAI three-month average both revise after first "
 "publication."
)
d.save('report_corrected.docx')
print('saved')
