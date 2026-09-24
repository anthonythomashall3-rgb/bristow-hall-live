import re
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE as RT
d = Document('report_corrected.docx')

def replace_text(p, old, new):
    runs=p.runs; full=''.join(r.text for r in runs); i=full.find(old)
    if i<0: return False
    j=i+len(old); pos=0; done=False
    for r in runs:
        a,b=pos,pos+len(r.text); pos=b
        if b<=i or a>=j: continue
        s=max(a,i)-a; e=min(b,j)-a
        if not done: r.text=r.text[:s]+new+r.text[e:]; done=True
        else: r.text=r.text[:s]+r.text[e:]
    return True

def plain(t):
    r=OxmlElement('w:r'); x=OxmlElement('w:t'); x.set(qn('xml:space'),'preserve'); x.text=t; r.append(x); return r

def add_links(p, urls):
    part=p.part
    for u in urls:
        rid=part.relate_to(u, RT.HYPERLINK, is_external=True)
        p._p.append(plain(' ')); p._p.append(plain('·')); p._p.append(plain(' '))
        hl=OxmlElement('w:hyperlink'); hl.set(qn('r:id'), rid)
        r=OxmlElement('w:r'); rPr=OxmlElement('w:rPr')
        st=OxmlElement('w:rStyle'); st.set(qn('w:val'),'Hyperlink'); rPr.append(st)
        un=OxmlElement('w:u'); un.set(qn('w:val'),'single'); rPr.append(un)
        r.append(rPr)
        x=OxmlElement('w:t'); x.set(qn('xml:space'),'preserve'); x.text=u
        r.append(x); hl.append(r); p._p.append(hl)

# --- punctuation: three ."". doubled periods
assert replace_text(d.paragraphs[79],
  'but a recession did not materialize.”.', 'but a recession did not materialize”.'), 'p79'
assert replace_text(d.paragraphs[225],
  'the role of the stance of monetary policy.”.', 'the role of the stance of monetary policy”.'), 'p225'
assert replace_text(d.paragraphs[307],
  'the Sahm rule recession indicator to states.”.', 'the Sahm rule recession indicator to states”.'), 'p307'

# --- citation completeness
assert replace_text(d.paragraphs[135],
  'Pilar Poncela (2018), International Journal of Forecasting;',
  'Pilar Poncela (2018), International Journal of Forecasting 34(4): 598–611;'), 'p135'

# --- links
ADD = {
 80: ['https://www.federalreserve.gov/econres/notes/feds-notes/assessing-recession-risks-with-state-level-data-20260107.html',
      'https://www.prnewswire.com/news-releases/us-consumer-confidence-rises-slightly-in-august-302231679.html'],
 136:['https://ideas.repec.org/a/eee/intfor/v34y2018i4p598-611.html'],
 226:['https://www.bostonfed.org/publications/current-policy-perspectives/2020/predicting-recessions-using-the-yield-curve.aspx'],
 302:['https://www.nber.org/news/business-cycle-dating-committee-announcement-june-8-2020',
      'https://www.nber.org/news/business-cycle-dating-committee-announcement-july-19-2021'],
 308:['https://www.bloomberg.com/opinion/articles/2024-03-27/stop-applying-the-sahm-rule-recession-indicator-to-states'],
 326:['https://www.cnbc.com/video/2022/09/16/the-u-s-will-experience-a-rolling-recession-says-ed-yardeni.html',
      'https://www.cnbc.com/video/2023/03/21/we-have-the-view-this-is-a-rolling-recession-says-charles-schwabs-liz-ann-sonders.html'],
 328:['https://www.prnewswire.com/news-releases/manufacturing-pmi-at-55-6-july-2026-ism-manufacturing-pmi-report-302840669.html'],
}
for idx,urls in ADD.items():
    p=d.paragraphs[idx]
    assert p.text.strip().startswith('http'), (idx,p.text[:60])
    add_links(p,urls)

d.save('report_corrected.docx')
print('saved')
