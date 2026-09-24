"""Suite 22 — the claims added or changed in the final holistic pass."""
import pandas as pd, re, sys
from docx import Document
D='fresh/fred/'
def s(n):
    import os
    p=D+n+'.csv'
    if not os.path.exists(p): p='/home/claude/sr/fred_'+n+'.csv'
    d=pd.read_csv(p); d.columns=['d','v']
    d['d']=pd.to_datetime(d.d); d['v']=pd.to_numeric(d.v,errors='coerce')
    return d.dropna().set_index('d')['v']
d=Document('Missed_Recession_Paper_v9.4_Full.docx')
T='\n'.join(p.text for p in d.paragraphs)
for tb in d.tables: T+='\n'+'\n'.join(c.text for r in tb.rows for c in r.cells)
T=re.sub(r'\s+',' ',T)
P=F=0
def chk(label, got, want):
    global P,F
    ok = (got==want)
    print(('PASS' if ok else 'FAIL'), label, '| got', got, '| want', want)
    P+=ok; F+=(not ok)
A=pd.Timestamp('2024-04-01'); B=pd.Timestamp('2024-08-01')
# the six indicators, April-August 2024
for k,exp in [('CMRMTSPL',1.87),('PCEC96',1.26),('W875RX1',0.78),('INDPRO',0.19),('PAYEMS',0.14),('CE16OV',-0.01)]:
    x=s(k); chk(f'{k} Apr-Aug 2024 pct', round(100*(x[B]-x[A])/x[A],2), exp)
chk('payroll change (thousands)', int(s('PAYEMS')[B]-s('PAYEMS')[A]), 227)
chk('household emp change (thousands)', int(s('CE16OV')[B]-s('CE16OV')[A]), -9)
chk('sales annualized', round(((s('CMRMTSPL')[B]/s('CMRMTSPL')[A])**3-1)*100,1), 5.7)
chk('PCE annualized', round(((s('PCEC96')[B]/s('PCEC96')[A])**3-1)*100,1), 3.8)
chk('IP annualized', round(((s('INDPRO')[B]/s('INDPRO')[A])**3-1)*100,1), 0.6)
# vacancy rate levels
jo=s('JTSJOL'); lf=s('CLF16OV'); vac=100*jo/lf
chk('vacancy rate Mar 2022 (openings/labor force)', round(float(vac[pd.Timestamp('2022-03-01')]),1), 7.5)
chk('vacancy rate Aug 2024 (openings/labor force)', round(float(vac[pd.Timestamp('2024-08-01')]),1), 4.5)
chk('vacancy rate Nov 2025 (openings/labor force)', round(float(vac[pd.Timestamp('2025-11-01')]),1), 4.0)
chk('vacancy fall Jul 2022 to Nov 2025, points', round(float(vac[pd.Timestamp('2022-07-01')]-vac[pd.Timestamp('2025-11-01')]),1), 3.1)
# never-fired instruments
chk('CFNAI-MA3 min since 2022', s('CFNAIMA3')[s('CFNAIMA3').index>='2022-01-01'].min(), -0.39)
chk('Chauvet-Piger max since 2022', s('RECPROUSM156N')[s('RECPROUSM156N').index>='2022-01-01'].max(), 2.3)
chk('Hamilton index max since 2022', s('JHGDPBRINDX')[s('JHGDPBRINDX').index>='2022-01-01'].max(), 37.4)
# Sahm crossing / peak on both vintages
for n,jul,aug in [('SAHMREALTIME',0.53,0.57),('SAHMCURRENT',0.50,0.57)]:
    x=s(n); chk(f'{n} Jul 2024', x[pd.Timestamp('2024-07-01')], jul)
    chk(f'{n} Aug 2024', x[pd.Timestamp('2024-08-01')], aug)
# prior-recession comparison
six=['CMRMTSPL','PCEC96','W875RX1','INDPRO','PAYEMS','CE16OV']
for lab,pk,exp_f,exp_t in [('1990-91','1990-07-01',5,5),('2001','2001-03-01',5,5),('2007-09','2007-12-01',5,6)]:
    Pk=pd.Timestamp(pk); E=Pk+pd.DateOffset(months=4); fell=tot=0
    for k in six:
        x=s(k)
        if Pk not in x.index or E not in x.index: continue
        tot+=1; fell += (x[E]<x[Pk])
    chk(f'{lab} first four months fell/total', (fell,tot), (exp_f,exp_t))
# text assertions added this pass
for frag in ['periods of economic contraction have historically been associated with values of the CFNAI-MA3 below',
             'the vacancy rate fell from 7.5 to 4.5 percent',
             'fell three points, to 4.0 percent by November 2025',
             '0.5 percent, in the Bureau',
             'reached exactly 30 in August 2024',
             '17 states in July, 30 in August',
             'the only month since March 2021 to reach the line']:
    chk(f'text present: {frag[:52]}', frag in T, True)
# and the removed wording is gone
for frag in ['a CFNAI-MA3 value below', 'came at or across that line in mid-2024', 'fell three percentage points']:
    chk(f'text removed: {frag[:44]}', frag in T, False)
# --- title-block centring and the acknowledgment footnote (added after the v9.1 fix)
import pdfplumber, warnings
warnings.filterwarnings('ignore')
pdf=pdfplumber.open('Missed_Recession_Paper_v9.4_Full.pdf')
pg=pdf.pages[0]
lines={}
for ch in pg.chars[:2500]:
    lines.setdefault(round(ch['top'],1),[]).append(ch)
rows=[]
for k in sorted(lines)[:6]:
    cs=lines[k]; x0=min(c['x0'] for c in cs); x1=max(c['x1'] for c in cs)
    rows.append((k, (x0+x1)/2, ''.join(c['text'] for c in sorted(cs,key=lambda c:c['x0']))))
def midof(sub):
    for k,mid,txt in rows:
        if sub in txt: return mid
    return None
for label,sub in [('title','Recession Signals'),('authors','Anthony Hall and Duke K. Bristow'),
                  ('affiliation','Marshall School of Business'),('working-paper line','FBE Working Paper')]:
    m=midof(sub)
    chk(f'{label} centered within 0.25pt of 306', (m is not None and abs(m-306)<=0.25), True)
alltext='\n'.join((q.extract_text() or '') for q in pdf.pages)
chk('acknowledgment footnote renders', 'Gioele Abbati' in alltext, True)
print(f'\nWITH TITLE-BLOCK CHECKS: PASSED {P}  FAILED {F}')
sys.exit(1 if F else 0)
