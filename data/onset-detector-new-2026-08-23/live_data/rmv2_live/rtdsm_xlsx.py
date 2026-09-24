"""B-LAND-4 §2 — RTDSM real-time-data xlsx matrix -> as-of observation cells.

The Philadelphia Fed Real-Time Data Set (RTDSM; Croushore-Stark, J. Econometrics
105, 2001) publishes ONE xlsx per variable.  Measured layout (research/AL4_layout.txt,
recorded before any parse-at-scale, per §2.1):

    row 1   DATE, <VAR><YY>Q<N>, <VAR><YY>Q<N>, ...   (quarterly-vintage columns)
        or  DATE, <VAR><YY>M<M>, <VAR><YY>M<M>, ...   (monthly-vintage columns)
    row 2+  <YYYY>:Q<N> | <YYYY>:<MM>, value, value, ...

Each column is one as-of snapshot (vintage); each row is one observation period;
each cell is the value known AS-OF that vintage.  ``#N/A`` means the observation
did not yet exist as-of that vintage -> absence, NOT a hole (skip, never zero).

This module is the ONE new counted shape (§2.2): pure, deterministic, no network,
no AI.  It emits raw ``(vintage, period, value)`` cells; series-id composition and
the family-naming decision live in ``series_id`` (own-base, never aliased), so the
naming ruling parameterises this module without touching the parse logic.

openpyxl 3.1.5 on py3.8 refuses the Phil-Fed workbooks' date-only ``<modified>``
core property; ``load_worksheet`` strips ``docProps/core.xml`` before loading
(the sheet bytes are untouched).
"""
from __future__ import annotations

import hashlib
import io
import re
import zipfile

import openpyxl

PARSER_VERSION = "rmv2-live/rtdsm_xlsx_matrix.v1"

# <VAR><YY>Q<N> (quarterly vintage) or <VAR><YY>M<M> (monthly vintage).
_VINTAGE_COL_RE = re.compile(r"^(?P<var>[A-Za-z0-9]+?)(?P<yy>\d{2})(?P<kind>[QM])(?P<n>\d{1,2})$")
# observation cell: YYYY:QN (quarterly obs) or YYYY:MM (monthly obs).
_OBS_Q_RE = re.compile(r"^(?P<y>\d{4}):Q(?P<q>[1-4])$")
_OBS_M_RE = re.compile(r"^(?P<y>\d{4}):(?P<m>\d{2})$")
# a landed value is a canonical decimal (optionally signed); anything else that
# is not the absence sentinel is a shape surprise and is refused, never coerced.
_DECIMAL_RE = re.compile(r"^-?\d+(?:\.\d+)?$")
_ABSENCE = frozenset(("#N/A", "", "."))
# quarter -> representative (middle) month of the quarter for the vintage stamp.
_Q_MIDMONTH = {"1": "02", "2": "05", "3": "08", "4": "11"}
# quarter -> first month for an OBSERVATION period (the observed quarter's start).
_Q_FIRSTMONTH = {"1": "01", "2": "04", "3": "07", "4": "10"}
# RTDSM vintages span 1962..present; 2-digit-year pivot at 50 covers 1950-2049.
_YEAR_PIVOT = 50


class RtdsmShapeError(ValueError):
    """A workbook, column, cell or value did not match the measured RTDSM shape.

    Raised rather than skipping, so a shape surprise can never silently
    understate what was landed (§19.4)."""


def _four_digit_year(yy: str) -> int:
    n = int(yy)
    return 2000 + n if n < _YEAR_PIVOT else 1900 + n


def vintage_stamp(colname: str, var: str) -> str:
    """``<VAR><YY>Q<N>``/``<VAR><YY>M<M>`` -> YYYYMMDD as-of stamp for ``var``.

    Quarterly vintages -> the quarter's middle month (02/05/08/11), day 01;
    monthly vintages -> that month, day 01.  Day is pinned to 01 (the panel
    column encodes only the vintage period, not a specific release day; the
    convention is documented rather than a release day being invented)."""
    m = _VINTAGE_COL_RE.match(str(colname))
    if not m or m.group("var") != var:
        raise RtdsmShapeError("column %r is not a %s vintage column" % (colname, var))
    year = _four_digit_year(m.group("yy"))
    kind, n = m.group("kind"), m.group("n")
    if kind == "Q":
        if n not in _Q_MIDMONTH:
            raise RtdsmShapeError("quarter out of range in %r" % colname)
        month = _Q_MIDMONTH[n]
    else:  # monthly
        mi = int(n)
        if not 1 <= mi <= 12:
            raise RtdsmShapeError("month out of range in %r" % colname)
        month = "%02d" % mi
    return "%04d%s01" % (year, month)


def obs_period(cell) -> str:
    """RTDSM DATE cell (``YYYY:QN`` or ``YYYY:MM``) -> canonical ``YYYY-MM-DD``.

    Quarterly obs -> first month of the quarter; monthly obs -> that month; day 01."""
    text = str(cell).strip()
    mq = _OBS_Q_RE.match(text)
    if mq:
        return "%s-%s-01" % (mq.group("y"), _Q_FIRSTMONTH[mq.group("q")])
    mm = _OBS_M_RE.match(text)
    if mm:
        month = int(mm.group("m"))
        if not 1 <= month <= 12:
            raise RtdsmShapeError("observation month out of range: %r" % cell)
        return "%s-%02d-01" % (mm.group("y"), month)
    raise RtdsmShapeError("unrecognised observation period: %r" % cell)


def series_id(family_prefix: str, var: str, stamp: str) -> str:
    """Own-base as-of series id: ``<PREFIX>_<VAR>.DEEPASOF<YYYYMMDD>``.

    Identity-neutral, per the owner's CLAIMSx/CMRMTSPLx ruling (B-LAND-3C-R2):
    each RTDSM variable is its OWN base; its mapping to a FRED concept
    (RUC~UNRATE, IPT/IPM~INDPRO, ...) is RESEARCH (§5), never an alias.  The
    ``<PREFIX>_`` guarantees the family prefix is disjoint from every FRED
    ALFRED base, so it cannot collide with a ``<BASE>.DEEPASOF`` FRED family."""
    return "%s_%s.DEEPASOF%s" % (family_prefix, var, stamp)


def load_worksheet(body: bytes):
    """First worksheet of an RTDSM workbook (docProps/core.xml stripped)."""
    src = zipfile.ZipFile(io.BytesIO(body))
    buf = io.BytesIO()
    dst = zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED)
    for item in src.infolist():
        if item.filename == "docProps/core.xml":
            continue
        dst.writestr(item, src.read(item.filename))
    dst.close()
    src.close()
    buf.seek(0)
    wb = openpyxl.load_workbook(buf, read_only=True, data_only=True)
    return wb[wb.sheetnames[0]], wb


def parse_matrix(body: bytes, var: str):
    """Parse one RTDSM variable workbook into as-of cells.

    Returns ``(cells, sha256, meta)`` where each cell is
    ``{"var", "vintage", "period", "value"}`` for every non-absence entry, sha256
    is over the ORIGINAL bytes (provenance), and meta carries the counted shape.
    Absence cells (``#N/A``) are skipped (not holes); any other malformed cell,
    column, or value raises ``RtdsmShapeError`` rather than being coerced/skipped."""
    sha = hashlib.sha256(body).hexdigest()
    ws, wb = load_worksheet(body)
    try:
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            raise RtdsmShapeError("workbook is empty")
        header = [None if c is None else str(c).strip() for c in header]
        if not header or header[0] != "DATE":
            raise RtdsmShapeError("header does not start with DATE: %r" % header[:2])
        vintage_cols = []  # (col_index, vintage_stamp)
        for idx, name in enumerate(header[1:], start=1):
            if name is None or name == "":
                continue
            vintage_cols.append((idx, vintage_stamp(name, var)))  # raises on foreign col
        if not vintage_cols:
            raise RtdsmShapeError("no vintage columns for %s" % var)

        cells = []
        seen = set()
        obs_rows = 0
        for raw in rows:
            if raw is None or raw[0] is None:
                continue
            period = obs_period(raw[0])
            obs_rows += 1
            for idx, stamp in vintage_cols:
                if idx >= len(raw):
                    continue
                cell = raw[idx]
                if cell is None:
                    continue
                text = str(cell).strip()
                if text in _ABSENCE:
                    continue
                if not _DECIMAL_RE.match(text):
                    raise RtdsmShapeError(
                        "non-decimal value %r at %s vintage-col#%d" % (text, period, idx)
                    )
                key = (stamp, period)
                if key in seen:
                    raise RtdsmShapeError("duplicate cell %s @ %s" % (stamp, period))
                seen.add(key)
                cells.append({"var": var, "vintage": stamp, "period": period, "value": text})
        meta = {
            "var": var,
            "vintages": len(vintage_cols),
            "obs_rows": obs_rows,
            "cells": len(cells),
            "parser_version": PARSER_VERSION,
        }
        if not cells:
            raise RtdsmShapeError("no landed observations for %s" % var)
        return cells, sha, meta
    finally:
        wb.close()
