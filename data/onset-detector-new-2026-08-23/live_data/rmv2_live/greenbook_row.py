"""Strict normalization for the Philadelphia Fed Greenbook/Tealbook Row Format xlsx.

B-ACQ-GREENBOOK. Real-time Fed staff readings of the economy, ONE row per Greenbook
publication (FOMC meeting). Each of the 15 variable worksheets carries the columns

    DATE, <VAR>B4, <VAR>B3, <VAR>B2, <VAR>B1, <VAR>F0 .. <VAR>F9, GBdate

where ``DATE`` is the anchor quarter (float ``YYYY.Q``), ``GBdate`` is the exact
publication date (int ``YYYYMMDD`` — the VINTAGE that makes the boundary provable),
the ``B4..B1`` columns are the real-time HISTORICAL values known at that Greenbook for
the 4/3/2/1 quarters preceding ``DATE``, and the ``F0..F9`` columns are the staff
PROJECTIONS (``F0`` = current-quarter nowcast, ``F1..F9`` = 1..9-quarter-ahead
forecasts).

Two DISJOINT output halves, NEVER mixed in one series (brief STOP; rulebook §3.1/§3.6):

  * ``hist`` -> ``GB_<VAR>_HIST`` : the ``Bx`` real-time historical values. Mode
    ``archive_snapshot_asof``. ``observation_period`` = quarter(``DATE`` - x); the
    dimension ``greenbook_vintage`` = the ``GBdate`` ISO publication date. No forecast
    horizon is ever attached to a HIST record.
  * ``proj`` -> ``GB_<VAR>_PROJ`` : the ``Fx`` staff projections. Mode
    ``substituted_diagnostic``. ``observation_period`` = quarter(``DATE`` + h) for
    ``F<h>``; ``forecast_origin`` = the ``GBdate`` ISO date; ``forecast_horizon`` = h.

``#N/A`` (and blank) cells emit ``value=None`` / ``value_status='unavailable'`` (the
FRED current-lane convention). Values are copied verbatim as the shortest
round-tripping decimal — never rounded, merged, spliced, or re-based. This module
lands data; it adopts nothing and classifies nothing against the member roster.
"""

from __future__ import absolute_import

import io

import openpyxl


class GreenbookRowDataError(ValueError):
    """Raised when the Greenbook Row Format xlsx bytes fail the reviewed contract."""


# The 15 variable worksheets, each -> (unit, human label root). Descriptions are
# copied verbatim from the workbook's own Documentation sheet (measured, not guessed).
GREENBOOK_VARIABLES = {
    "gRGDP": ("Annualized percentage points",
              "Q/Q growth in real GDP, chain weight"),
    "gPGDP": ("Annualized percentage points",
              "Q/Q growth in price index for GDP, chain weight"),
    "UNEMP": ("Percentage points", "unemployment rate"),
    "gPCPI": ("Annualized percentage points", "Q/Q headline CPI inflation"),
    "gPCPIX": ("Annualized percentage points", "Q/Q core CPI inflation"),
    "gPPCE": ("Annualized percentage points",
              "Q/Q headline PCE inflation, chain weight"),
    "gPPCEX": ("Annualized percentage points",
               "Q/Q core PCE inflation, chain weight"),
    "gRPCE": ("Annualized percentage points",
              "Q/Q growth in real personal consumption expenditures, chain weight"),
    "gRBF": ("Annualized percentage points",
             "Q/Q growth in real business fixed investment, chain weight"),
    "gRRES": ("Annualized percentage points",
              "Q/Q growth in real residential investment, chain weight"),
    "gRGOVF": ("Annualized percentage points",
               "Q/Q growth in real federal government consumption & gross investment"),
    "gRGOVSL": ("Annualized percentage points",
                "Q/Q growth in real state & local government consumption & gross "
                "investment"),
    "gNGDP": ("Annualized percentage points", "Q/Q growth in nominal GDP"),
    "HSTART": ("Millions of units", "housing starts"),
    "gIP": ("Annualized percentage points",
            "Q/Q growth in the industrial production index"),
}

_BACKCAST_OFFSETS = (("B4", -4), ("B3", -3), ("B2", -2), ("B1", -1))
_MAX_FORECAST_HORIZON = 9  # F0..F9
_MISSING_TOKENS = frozenset(("", "#N/A", "#n/a", "NA", "N/A", "#VALUE!"))


def _quarter_period(year, quarter):
    """Map a (year, quarter) to the canonical first-of-quarter ISO date."""
    return "%04d-%02d-01" % (year, quarter * 3 - 2)


def _shift_quarter(year, quarter, delta):
    """Shift (year, quarter in 1..4) by ``delta`` quarters, returning (year, quarter)."""
    ordinal = year * 4 + (quarter - 1) + delta
    new_year, new_index = divmod(ordinal, 4)
    return new_year, new_index + 1


def _parse_anchor(value):
    """``1967.1`` -> (1967, 1). Rejects anything outside quarters 1-4 / 1900-2100."""
    try:
        number = float(value)
    except (TypeError, ValueError):
        raise GreenbookRowDataError("greenbook DATE was non-numeric: %r" % (value,))
    year = int(number)
    quarter = int(round((number - year) * 10))
    if quarter < 1 or quarter > 4:
        raise GreenbookRowDataError("greenbook DATE quarter out of range: %r" % (value,))
    if year < 1900 or year > 2100:
        raise GreenbookRowDataError("greenbook DATE year out of range: %r" % (value,))
    return year, quarter


def _parse_gbdate(value):
    """``19670329`` (int/str) -> ``'1967-03-29'`` with a real calendar-date check."""
    try:
        digits = "%08d" % int(value)
    except (TypeError, ValueError):
        raise GreenbookRowDataError("greenbook GBdate was non-numeric: %r" % (value,))
    if len(digits) != 8:
        raise GreenbookRowDataError("greenbook GBdate not YYYYMMDD: %r" % (value,))
    year, month, day = int(digits[:4]), int(digits[4:6]), int(digits[6:8])
    try:
        import datetime
        datetime.date(year, month, day)
    except ValueError as exc:
        raise GreenbookRowDataError("greenbook GBdate not a real date: %r (%s)"
                                    % (value, exc))
    return "%04d-%02d-%02d" % (year, month, day)


def _canonical_value(cell):
    """Return (value, status): shortest round-tripping decimal for a real number, or
    (None, 'unavailable') for a missing cell. NEVER rounds."""
    if cell is None:
        return None, "unavailable"
    if isinstance(cell, str):
        token = cell.strip()
        if token in _MISSING_TOKENS:
            return None, "unavailable"
        try:
            number = float(token)
        except ValueError:
            raise GreenbookRowDataError("greenbook cell was not numeric: %r" % (cell,))
    elif isinstance(cell, bool):
        raise GreenbookRowDataError("greenbook cell was a bool: %r" % (cell,))
    elif isinstance(cell, (int, float)):
        number = float(cell)
    else:
        raise GreenbookRowDataError("greenbook cell had unexpected type: %r" % (cell,))
    if number != number or number in (float("inf"), float("-inf")):
        raise GreenbookRowDataError("greenbook cell not finite: %r" % (cell,))
    rendered = repr(number)
    return rendered, "actual"


def _load_workbook(body):
    if not isinstance(body, (bytes, bytearray)) or not body:
        raise GreenbookRowDataError("greenbook xlsx body must be nonempty bytes")
    try:
        return openpyxl.load_workbook(io.BytesIO(bytes(body)), read_only=True,
                                      data_only=True)
    except Exception as exc:  # openpyxl raises a spread of types on bad input
        raise GreenbookRowDataError("openpyxl could not load greenbook xlsx: %s" % exc)


def _header_index(header, name):
    for index, cell in enumerate(header):
        if isinstance(cell, str) and cell.strip() == name:
            return index
    raise GreenbookRowDataError("greenbook sheet missing column %r" % (name,))


def parse_greenbook_row_workbook(series_config, body):
    """Parse the Greenbook Row Format workbook into one half's observations.

    ``series_config`` (the source's ``series`` object) must declare
    ``{"half": "hist"|"proj"}``. Returns a list of observation dicts; the adapter
    wrapper turns each into a store record.
    """
    if not isinstance(series_config, dict):
        raise GreenbookRowDataError("greenbook series config must be an object")
    half = series_config.get("half")
    if half not in ("hist", "proj"):
        raise GreenbookRowDataError("greenbook half must be 'hist' or 'proj': %r"
                                    % (half,))
    workbook = _load_workbook(body)
    try:
        present = set(workbook.sheetnames)
        missing = [v for v in GREENBOOK_VARIABLES if v not in present]
        if missing:
            raise GreenbookRowDataError(
                "greenbook workbook missing variable sheets: %r" % (missing,))
        observations = []
        for var, (unit, label_root) in GREENBOOK_VARIABLES.items():
            worksheet = workbook[var]
            rows = worksheet.iter_rows(values_only=True)
            header = list(next(rows))
            date_i = _header_index(header, "DATE")
            gb_i = _header_index(header, "GBdate")
            if half == "hist":
                col_idx = [(_header_index(header, var + suf), delta)
                           for suf, delta in _BACKCAST_OFFSETS]
                series_id = "GB_%s_HIST" % var
            else:
                col_idx = [(_header_index(header, "%sF%d" % (var, h)), h)
                           for h in range(_MAX_FORECAST_HORIZON + 1)]
                series_id = "GB_%s_PROJ" % var
            for row in rows:
                if row[date_i] is None and row[gb_i] is None:
                    continue
                year, quarter = _parse_anchor(row[date_i])
                vintage = _parse_gbdate(row[gb_i])
                for col, delta in col_idx:
                    tgt_year, tgt_q = _shift_quarter(year, quarter, delta)
                    period = _quarter_period(tgt_year, tgt_q)
                    value, status = _canonical_value(row[col])
                    obs = {
                        "series_id": series_id,
                        "observation_period": period,
                        "value": value,
                        "value_status": status,
                        "unit": unit,
                        "greenbook_vintage": vintage,
                        "variable_code": var,
                    }
                    if half == "hist":
                        obs["label"] = ("Greenbook real-time historical %s [as known "
                                        "at %s]" % (label_root, vintage))
                    else:
                        obs["label"] = ("Greenbook staff projection of %s [origin %s, "
                                        "h=%d]" % (label_root, vintage, delta))
                        obs["forecast_origin"] = vintage
                        obs["forecast_horizon"] = delta
                    observations.append(obs)
        observations.sort(key=lambda o: (o["series_id"], o["greenbook_vintage"],
                                         o["observation_period"]))
        return observations
    finally:
        workbook.close()
