"""B-LAND-8 — Philadelphia Fed SPF / Anxious Index xlsx adapter.

FORECAST class. openpyxl chokes on the SAS-written docProps/core.xml (non-zero-padded
W3CDTF hour); the adapter repairs it IN MEMORY and the binder stores the ORIGINAL
cache bytes (CH-R49 openpyxl_core_defect). Counts are measured from the cache bytes
(research/bland8_truth.py). No network, deterministic.
"""
import hashlib
import os
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from live_data.rmv2_live.adapters import normalize  # noqa: E402
from live_data.rmv2_live.forecast_xlsx import (  # noqa: E402
    ForecastXlsxDataError,
    parse_forecast_xlsx_bytes,
    repair_core_xml_bytes,
)

CACHE = os.path.join(str(PROJECT_ROOT), "research", "prefetch", "forecasts")
ANXIOUS = os.path.join(CACHE, "anxious_index", "anxious_index_chart.xlsx")
MEAN_RGDP_GROWTH = os.path.join(CACHE, "spf", "mean_rgdp_growth.xlsx")
MEDIAN_UNEMP_LEVEL = os.path.join(CACHE, "spf", "median_unemp_level.xlsx")
INDIV_CORECPI = os.path.join(CACHE, "spf", "individual_corecpi.xlsx")
INDIV_RGDP = os.path.join(CACHE, "spf", "individual_rgdp.xlsx")


def _body(path):
    with open(path, "rb") as handle:
        return handle.read()


def _anxious_config():
    return {
        "forecast_shape": "anxious_index",
        "sheet": "Data",
        "series_id": "SPF_ANXIOUS_INDEX",
        "unit": "percent probability",
        "label": "Anxious Index (SPF prob of negative q/q real GDP growth)",
    }


def _mean_growth_config():
    return {
        "forecast_shape": "spf_aggregate",
        "sheet": "Mean_Growth",
        "statistic": "MEAN",
        "unit": "SPF mean growth forecast",
        "label": "SPF mean growth forecast",
    }


def _median_level_config():
    return {
        "forecast_shape": "spf_aggregate",
        "sheet": "Median_Level",
        "statistic": "MEDIAN",
        "unit": "SPF median level forecast",
        "label": "SPF median level forecast",
    }


def _individual_config(var="CORECPI", sheet=None):
    return {
        "forecast_shape": "spf_individual",
        "sheet": sheet or var,
        "spf_target_code": var,
        "id_prefix": "SPF_",
        "unit": "SPF individual-forecaster point forecast",
        "label": "SPF individual-forecaster panel",
    }


# ---- core.xml repair -------------------------------------------------------

def test_raw_bytes_fail_openpyxl_but_repair_unlocks_them():
    import openpyxl
    body = _body(MEAN_RGDP_GROWTH)
    with pytest.raises(Exception):
        openpyxl.load_workbook(__import__("io").BytesIO(body), read_only=True)
    # repaired handle loads
    wb = openpyxl.load_workbook(repair_core_xml_bytes(body), read_only=True)
    assert "Mean_Growth" in wb.sheetnames
    wb.close()


def test_repair_does_not_mutate_original_bytes():
    body = _body(MEAN_RGDP_GROWTH)
    before = hashlib.sha256(body).hexdigest()
    repair_core_xml_bytes(body)
    assert hashlib.sha256(body).hexdigest() == before


def test_repair_rejects_non_zip():
    with pytest.raises(ForecastXlsxDataError):
        repair_core_xml_bytes(b"not a zip")


def test_hour_pad_regex_forms():
    from live_data.rmv2_live.forecast_xlsx import _CORE_XML_HOUR_RE
    assert _CORE_XML_HOUR_RE.sub(rb"T0\1:", b"2026-05-13T 2:31:35") == b"2026-05-13T02:31:35"
    assert _CORE_XML_HOUR_RE.sub(rb"T0\1:", b"2026-05-13T2:31:35") == b"2026-05-13T02:31:35"
    # a well-formed two-digit hour is left untouched
    assert _CORE_XML_HOUR_RE.sub(rb"T0\1:", b"2026-05-13T12:31:35") == b"2026-05-13T12:31:35"


# ---- Anxious Index ---------------------------------------------------------

def test_anxious_counts_and_endpoints():
    obs = parse_forecast_xlsx_bytes(_anxious_config(), _body(ANXIOUS))
    assert all(o["series_id"] == "SPF_ANXIOUS_INDEX" for o in obs)
    assert all(o["value_status"] == "actual" for o in obs)  # future quarters skipped
    assert len(obs) == 231
    assert obs[0]["observation_period"] == "1969-01-01"
    assert obs[0]["value"] == "17.69"
    assert obs[-1]["observation_period"] == "2026-07-01"
    assert obs[-1]["value"] == "25.0542"
    assert obs[0]["spf_target_code"] == "ANXIOUS_INDEX"


def test_anxious_periods_are_unique_and_sorted():
    obs = parse_forecast_xlsx_bytes(_anxious_config(), _body(ANXIOUS))
    periods = [o["observation_period"] for o in obs]
    assert periods == sorted(periods)
    assert len(periods) == len(set(periods))


# ---- SPF aggregate ---------------------------------------------------------

def test_mean_growth_column_becomes_own_series():
    obs = parse_forecast_xlsx_bytes(_mean_growth_config(), _body(MEAN_RGDP_GROWTH))
    ids = sorted(set(o["series_id"] for o in obs))
    assert ids == [
        "SPF_DRGDP2_MEAN", "SPF_DRGDP3_MEAN", "SPF_DRGDP4_MEAN",
        "SPF_DRGDP5_MEAN", "SPF_DRGDP6_MEAN",
    ]
    d2 = [o for o in obs if o["series_id"] == "SPF_DRGDP2_MEAN"]
    assert len(d2) == 231
    assert all(o["value_status"] == "actual" for o in d2)
    assert d2[0]["observation_period"] == "1968-10-01"
    assert d2[0]["value"] == "2.0398"
    assert d2[-1]["observation_period"] == "2026-04-01"


def test_mean_growth_missing_cells_are_unavailable_not_dropped():
    obs = parse_forecast_xlsx_bytes(_mean_growth_config(), _body(MEAN_RGDP_GROWTH))
    d6 = [o for o in obs if o["series_id"] == "SPF_DRGDP6_MEAN"]
    # DRGDP6: 231 rows, 226 available, 5 unavailable (measured)
    assert len(d6) == 231
    assert sum(1 for o in d6 if o["value_status"] == "actual") == 226
    assert sum(1 for o in d6 if o["value_status"] == "unavailable") == 5
    assert all(o["value"] is None for o in d6 if o["value_status"] == "unavailable")


def test_median_level_statistic_suffix():
    obs = parse_forecast_xlsx_bytes(_median_level_config(), _body(MEDIAN_UNEMP_LEVEL))
    assert all(o["series_id"].endswith("_MEDIAN") for o in obs)
    assert "SPF_UNEMP1_MEDIAN" in {o["series_id"] for o in obs}


def test_deterministic_across_two_parses():
    a = parse_forecast_xlsx_bytes(_mean_growth_config(), _body(MEAN_RGDP_GROWTH))
    b = parse_forecast_xlsx_bytes(_mean_growth_config(), _body(MEAN_RGDP_GROWTH))
    assert a == b


# ---- SPF individual panels (B-LAND-8-R2) -----------------------------------

def test_individual_horizon_becomes_own_series():
    obs = parse_forecast_xlsx_bytes(_individual_config("CORECPI"), _body(INDIV_CORECPI))
    ids = sorted(set(o["series_id"] for o in obs))
    # CORECPI has 9 horizon columns (1-6, A-C); each becomes SPF_CORECPI_IND_<H>.
    assert ids == [
        "SPF_CORECPI_IND_1", "SPF_CORECPI_IND_2", "SPF_CORECPI_IND_3",
        "SPF_CORECPI_IND_4", "SPF_CORECPI_IND_5", "SPF_CORECPI_IND_6",
        "SPF_CORECPI_IND_A", "SPF_CORECPI_IND_B", "SPF_CORECPI_IND_C",
    ]


def test_individual_carries_panel_dimensions():
    obs = parse_forecast_xlsx_bytes(_individual_config("CORECPI"), _body(INDIV_CORECPI))
    assert obs  # non-empty
    for o in obs:
        assert o["spf_target_code"] == "CORECPI"
        assert isinstance(o["spf_forecaster_id"], int)
        assert o["spf_horizon"] in {"1", "2", "3", "4", "5", "6", "A", "B", "C"}
        # industry is an int code (1-3) or None when the file wrote #N/A.
        assert o["spf_industry"] is None or isinstance(o["spf_industry"], int)


def test_individual_is_sparse_actual_only():
    # A forecaster x quarter x horizon grid is >99% empty; the panel emits ONLY
    # the actual (non-#N/A) cells, never a dense unavailable row per cell.
    obs = parse_forecast_xlsx_bytes(_individual_config("CORECPI"), _body(INDIV_CORECPI))
    assert all(o["value_status"] == "actual" for o in obs)
    assert all(o["value"] is not None for o in obs)


def test_individual_panel_key_is_unique():
    obs = parse_forecast_xlsx_bytes(_individual_config("RGDP"), _body(INDIV_RGDP))
    keys = [
        (o["series_id"], o["observation_period"], o["spf_forecaster_id"],
         o["spf_industry"])
        for o in obs
    ]
    assert len(keys) == len(set(keys))


def test_individual_deterministic_across_two_parses():
    a = parse_forecast_xlsx_bytes(_individual_config("CORECPI"), _body(INDIV_CORECPI))
    b = parse_forecast_xlsx_bytes(_individual_config("CORECPI"), _body(INDIV_CORECPI))
    assert a == b


def test_individual_rejects_missing_target_code():
    cfg = _individual_config("CORECPI")
    del cfg["spf_target_code"]
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(cfg, _body(INDIV_CORECPI))


def test_individual_rejects_target_not_matching_columns():
    # spf_target_code must be the file's varcode; a mismatch means the wrong
    # file was bound to this identity — refuse rather than mislabel.
    cfg = _individual_config("WRONGVAR")
    cfg["sheet"] = "CORECPI"
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(cfg, _body(INDIV_CORECPI))


def test_individual_series_ids_are_spf_namespaced():
    obs = parse_forecast_xlsx_bytes(_individual_config("RGDP"), _body(INDIV_RGDP))
    assert all(o["series_id"].startswith("SPF_") for o in obs)
    assert all("_IND_" in o["series_id"] for o in obs)


def test_individual_feed_factory_identity_shape():
    from live_data.rmv2_live.feed_factory import _adapter_output_series_ids
    src = {
        "adapter": "forecast_xlsx",
        "series": {
            "forecast_shape": "spf_individual",
            "spf_target_code": "CORECPI",
            "id_prefix": "SPF_",
            "sheet": "CORECPI",
        },
    }
    # data-dependent ids: derives an empty exact set (family gate governs).
    assert _adapter_output_series_ids(src, require_complete=True) == frozenset()


def test_individual_feed_factory_rejects_missing_target():
    from live_data.rmv2_live.feed_factory import _adapter_output_series_ids
    from live_data.rmv2_live.canonical import CanonicalDataError
    src = {
        "adapter": "forecast_xlsx",
        "series": {
            "forecast_shape": "spf_individual",
            "id_prefix": "SPF_",
            "sheet": "CORECPI",
        },
    }
    with pytest.raises(CanonicalDataError):
        _adapter_output_series_ids(src, require_complete=True)


def test_individual_normalize_carries_dims_and_source():
    src = {
        "adapter": "forecast_xlsx",
        "endpoint": "https://www.philadelphiafed.org/spf",
        "information_set_mode": "substituted_diagnostic",
        "method_version": "spf_individual_forecast_offline.v1",
        "publisher": "Federal Reserve Bank of Philadelphia",
        "publisher_release_clock": "quarterly named survey vintage",
        "rights_status": "published_survey_data",
        "series": _individual_config("CORECPI"),
        "source_id": "philadelphia_spf_individual_corecpi",
        "value_status": "actual",
    }
    records = normalize(src, _body(INDIV_CORECPI), "2026-08-06T00:00:00Z")
    assert records
    for record in records:
        assert record["source_id"] == "philadelphia_spf_individual_corecpi"
        assert record["information_set_mode"] == "substituted_diagnostic"
        assert record["series_id"].startswith("SPF_CORECPI_IND_")
        assert record["spf_target_code"] == "CORECPI"
        assert isinstance(record["spf_forecaster_id"], int)
        assert record["spf_horizon"] in {"1", "2", "3", "4", "5", "6",
                                         "A", "B", "C"}


# ---- guards ----------------------------------------------------------------

def test_rejects_unknown_shape():
    cfg = _anxious_config()
    cfg["forecast_shape"] = "bogus"
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(cfg, _body(ANXIOUS))


def test_rejects_non_spf_series_id():
    cfg = _anxious_config()
    cfg["series_id"] = "ICSA"  # a modern twin id must be refused
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(cfg, _body(ANXIOUS))


def test_rejects_bad_statistic():
    cfg = _mean_growth_config()
    cfg["statistic"] = "AVG"
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(cfg, _body(MEAN_RGDP_GROWTH))


def test_rejects_empty_body():
    with pytest.raises(ForecastXlsxDataError):
        parse_forecast_xlsx_bytes(_anxious_config(), b"")


# ---- normalize integration -------------------------------------------------

def test_normalize_stamps_source_and_mode():
    src = {
        "adapter": "forecast_xlsx",
        "endpoint": "https://www.philadelphiafed.org/spf",
        "information_set_mode": "substituted_diagnostic",
        "method_version": "spf_forecast_offline.v1",
        "publisher": "Federal Reserve Bank of Philadelphia",
        "publisher_release_clock": "quarterly named survey vintage",
        "rights_status": "published_survey_data",
        "series": _mean_growth_config(),
        "source_id": "philadelphia_spf_mean_rgdp_growth",
        "value_status": "actual",
    }
    records = normalize(src, _body(MEAN_RGDP_GROWTH), "2026-08-06T00:00:00Z")
    assert records
    for record in records:
        assert record["information_set_mode"] == "substituted_diagnostic"
        assert record["source_id"] == "philadelphia_spf_mean_rgdp_growth"
        assert record["series_id"].startswith("SPF_DRGDP")
        assert record["value_status"] in ("actual", "unavailable")
        assert record["spf_target_code"].startswith("DRGDP")
