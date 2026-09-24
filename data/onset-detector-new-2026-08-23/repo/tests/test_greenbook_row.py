"""B-ACQ-GREENBOOK — Philadelphia Fed Greenbook Row Format xlsx adapter.

Real-time Fed staff readings, one row per Greenbook publication. The adapter splits
each variable into GB_<VAR>_HIST (Bx real-time historical, archive_snapshot_asof) and
GB_<VAR>_PROJ (Fx staff projections, substituted_diagnostic), keyed by the GBdate
publication vintage. The two halves are NEVER mixed. Synthetic fixture with
independently-computed expected values; a real-file smoke test runs when the staged
Row Format workbook is present. No network, deterministic.
"""
import io
import os
import sys
from pathlib import Path

import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import openpyxl  # noqa: E402

from live_data.rmv2_live.greenbook_row import (  # noqa: E402
    GREENBOOK_VARIABLES,
    GreenbookRowDataError,
    parse_greenbook_row_workbook,
    _shift_quarter,
    _parse_anchor,
    _parse_gbdate,
)
from live_data.rmv2_live.adapters import normalize  # noqa: E402
from live_data.rmv2_live.pipeline import _project_source_snapshot  # noqa: E402
from live_data.rmv2_live.canonical import CanonicalDataError  # noqa: E402

STAGED = os.path.join(
    str(PROJECT_ROOT), "research", "_staging", "greenbook", "greenbook-data",
    "documentation", "GBweb_Row_Format.xlsx")


def _synthetic_workbook():
    """One data row per variable sheet, schema DATE,B4..B1,F0..F9,GBdate.

    Row: DATE=1967.1 (=1967Q1), GBdate=19670329. Bx = 10+x, Fx = 20+h, with a single
    '#N/A' at F9 to exercise the missing path.
    """
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for var in GREENBOOK_VARIABLES:
        ws = wb.create_sheet(var)
        header = ["DATE"] + ["%sB%d" % (var, b) for b in (4, 3, 2, 1)] \
            + ["%sF%d" % (var, h) for h in range(10)] + ["GBdate"]
        ws.append(header)
        row = [1967.1] + [10 + b for b in (4, 3, 2, 1)] \
            + [20 + h for h in range(9)] + ["#N/A", 19670329]
        ws.append(row)
    doc = wb.create_sheet("Documentation", 0)
    doc.append(["Documentation"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_quarter_helpers():
    assert _parse_anchor(1967.1) == (1967, 1)
    assert _parse_anchor(2020.4) == (2020, 4)
    assert _shift_quarter(1967, 1, -1) == (1966, 4)
    assert _shift_quarter(1967, 1, -4) == (1966, 1)
    assert _shift_quarter(1967, 1, 4) == (1968, 1)
    assert _parse_gbdate(19670329) == "1967-03-29"
    with pytest.raises(GreenbookRowDataError):
        _parse_gbdate(19670231)  # Feb 31 is not a real date
    with pytest.raises(GreenbookRowDataError):
        _parse_anchor(1967.5)  # quarter 5 out of range


def test_hist_half_semantics():
    body = _synthetic_workbook()
    obs = parse_greenbook_row_workbook({"half": "hist"}, body)
    # 15 vars * 4 backcast columns = 60 HIST records
    assert len(obs) == 15 * 4
    assert {o["series_id"] for o in obs} == {"GB_%s_HIST" % v for v in GREENBOOK_VARIABLES}
    # firewall: no HIST record carries a forecast horizon
    assert all("forecast_horizon" not in o for o in obs)
    unemp = {o["observation_period"]: o for o in obs if o["series_id"] == "GB_UNEMP_HIST"}
    # DATE=1967Q1: B1 -> 1966Q4 (=1966-10-01) value 11; B4 -> 1966Q1 value 14
    assert unemp["1966-10-01"]["value"] == "11.0"
    assert unemp["1966-10-01"]["greenbook_vintage"] == "1967-03-29"
    assert unemp["1966-01-01"]["value"] == "14.0"
    # every HIST period precedes the publication-quarter
    for o in obs:
        assert o["observation_period"] < "1967-01-01"


def test_proj_half_semantics():
    body = _synthetic_workbook()
    obs = parse_greenbook_row_workbook({"half": "proj"}, body)
    # 15 vars * 10 forecast columns = 150 PROJ records
    assert len(obs) == 15 * 10
    assert {o["series_id"] for o in obs} == {"GB_%s_PROJ" % v for v in GREENBOOK_VARIABLES}
    unemp = [o for o in obs if o["series_id"] == "GB_UNEMP_PROJ"]
    by_h = {o["forecast_horizon"]: o for o in unemp}
    assert set(by_h) == set(range(10))
    # F0 -> nowcast for 1967Q1 (=1967-01-01), value 20, origin the GBdate
    assert by_h[0]["observation_period"] == "1967-01-01"
    assert by_h[0]["value"] == "20.0"
    assert by_h[0]["forecast_origin"] == "1967-03-29"
    # F4 -> 1968Q1
    assert by_h[4]["observation_period"] == "1968-01-01"
    # F9 was '#N/A' -> unavailable
    assert by_h[9]["value"] is None
    assert by_h[9]["value_status"] == "unavailable"


def test_halves_are_disjoint():
    body = _synthetic_workbook()
    hist = {o["series_id"] for o in parse_greenbook_row_workbook({"half": "hist"}, body)}
    proj = {o["series_id"] for o in parse_greenbook_row_workbook({"half": "proj"}, body)}
    assert hist.isdisjoint(proj)


def test_bad_half_rejected():
    with pytest.raises(GreenbookRowDataError):
        parse_greenbook_row_workbook({"half": "both"}, _synthetic_workbook())


def test_normalize_adapter_records():
    body = _synthetic_workbook()
    source = {
        "adapter": "greenbook_row_xlsx",
        "source_id": "greenbook_row_hist_offline",
        "endpoint": "https://www.philadelphiafed.org/greenbook",
        "information_set_mode": "archive_snapshot_asof",
        "method_version": "greenbook.v1",
        "publisher_release_clock": None,
        "rights_status": "public_source_reuse_terms_and_attribution_review_required",
        "value_status": "actual",
        "series": {"half": "hist"},
    }
    records = normalize(source, body, "2026-08-08T00:00:00Z")
    assert len(records) == 15 * 4
    r = records[0]
    assert r["information_set_mode"] == "archive_snapshot_asof"
    assert r["greenbook_vintage"] == "1967-03-29"
    assert r["provider_vintage_kind"] == "greenbook_realtime_historical"
    assert r["forecast_horizon"] is None  # _base_record default, HIST never sets it


def _greenbook_source(half, mode):
    return {
        "adapter": "greenbook_row_xlsx",
        "source_id": "greenbook_row_%s_offline" % half,
        "endpoint": "https://www.philadelphiafed.org/greenbook",
        "information_set_mode": mode,
        "method_version": "greenbook.v1",
        "publisher_release_clock": None,
        "rights_status": "published_research_data",
        "value_status": "actual",
        "series": {"half": half},
    }


def test_vintage_panel_projection_excludes_flat_series():
    body = _synthetic_workbook()
    # duplicate observation periods across vintages is inherent; the panel policy
    # keeps the records as evidence but emits NO flat series (excluded_from_flat).
    for half, mode in (("hist", "archive_snapshot_asof"),
                       ("proj", "substituted_diagnostic")):
        recs = normalize(_greenbook_source(half, mode), body, "2026-08-08T00:00:00Z")
        proj = _project_source_snapshot(
            "greenbook_row_%s_offline" % half, {"b": 1}, recs,
            "greenbook_vintage_panel.v1")
        assert proj["series"] == {}
        assert proj["record_count"] == len(recs)


def test_vintage_panel_rejects_missing_vintage():
    body = _synthetic_workbook()
    recs = normalize(_greenbook_source("hist", "archive_snapshot_asof"), body,
                     "2026-08-08T00:00:00Z")
    recs[0] = dict(recs[0])
    recs[0].pop("greenbook_vintage", None)
    with pytest.raises(CanonicalDataError):
        _project_source_snapshot("greenbook_row_hist_offline", {}, recs,
                                 "greenbook_vintage_panel.v1")


def test_vintage_panel_rejects_duplicate_key():
    body = _synthetic_workbook()
    recs = normalize(_greenbook_source("hist", "archive_snapshot_asof"), body,
                     "2026-08-08T00:00:00Z")
    dup = dict(recs[0])  # identical (series_id, period, vintage, horizon)
    with pytest.raises(CanonicalDataError):
        _project_source_snapshot("greenbook_row_hist_offline", {}, recs + [dup],
                                 "greenbook_vintage_panel.v1")


@pytest.mark.skipif(not os.path.exists(STAGED), reason="staged Row Format not present")
def test_real_workbook_boundary():
    with open(STAGED, "rb") as fh:
        body = fh.read()
    hist = parse_greenbook_row_workbook({"half": "hist"}, body)
    proj = parse_greenbook_row_workbook({"half": "proj"}, body)
    assert len(hist) > 0 and len(proj) > 0
    vintages = sorted({o["greenbook_vintage"] for o in hist})
    # measured span from the bytes (do not take 1966-2020 on faith, brief step 3)
    assert vintages[0] == "1967-03-29"
    assert vintages[-1] == "2020-12-04"
    assert {o["series_id"] for o in hist} == {"GB_%s_HIST" % v for v in GREENBOOK_VARIABLES}
