import openpyxl, json, os, time
from openpyxl.reader.excel import ExcelReader
ExcelReader.read_properties=lambda self: None   # skip buggy core.xml date parse
rows=json.load(open("research/_ch_r39_parsed.json"))
base="research/prefetch/rtdsm"
t0=time.time()
def parse(xf):
    wb=openpyxl.load_workbook(xf, read_only=True, data_only=True)
    ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True)
    hdr=next(it)
    vcols=[str(h) for h in hdr[1:] if h not in (None,"")]
    nv=len(vcols)
    dfirst=dlast=None; recs=0
    for r in it:
        d=r[0]
        if d in (None,""): continue
        if dfirst is None: dfirst=str(d)
        dlast=str(d)
        for c in r[1:1+nv]:
            if c is None: continue
            s=str(c).strip()
            if s and s!="#N/A" and s.upper()!="NA": recs+=1
    wb.close()
    return dict(vintages=nv,first_vintage=vcols[0] if vcols else None,
        last_vintage=vcols[-1] if vcols else None,obs_first=dfirst,obs_last=dlast,records=recs)
for r in rows:
    if "error" in r:
        try:
            r.update(parse(f"{base}/{r['var']}/{r['file']}")); r.pop("error")
        except Exception as e:
            r["error"]=str(e)[:80]
json.dump(rows, open("research/_ch_r39_parsed.json","w"), indent=0)
print("remaining errors:",sum(1 for r in rows if "error" in r),"elapsed",round(time.time()-t0,1))
