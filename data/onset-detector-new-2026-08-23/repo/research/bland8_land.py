"""B-LAND-8 — land the Anxious Index + the full SPF mean/median level/growth
AGGREGATE tranche as FORECAST-class offline-current EVIDENCE lanes (one source per
cached xlsx file), via the generic offline-current binder.

Owner ruling (2026-08-06, answer on disk for the B-LAND-7 OWNER-DECISION): land the
FULL SPF tranche. This script lands items 1 (Anxious) + 2 (aggregates). The 21
individual-forecaster panels (item 3) are a batch-sanctioned SEPARATE stop — they
need a store panel model (forecaster_id x industry x horizon, duplicate survey-quarter
periods) that this offline-current lane does not yet express. See the brief.

FORECAST FIREWALL (CLAUDE.md): forecasts of MEASUREMENTS are eligible LABELED
substituted_diagnostic inputs AND comparators; the RECESS/Anxious probabilities are
GDP-decline probabilities, NEVER NBER recession-LABEL products. Bound to the registered
forecast family `philadelphia_spf` (role forecast_comparator, access B).

Zero network. The cache xlsx bytes ARE the source object (sha-manifested from CH-R35),
normalized through the new `forecast_xlsx` adapter (in-memory core.xml repair; ORIGINAL
bytes bound). enabled:false + archival:true (frozen forecast-track snapshot). Converge
via ONE TARGETED zero-network refresh of the newly-bound DISABLED source_ids only (never
a full pipeline.refresh — the B-OFFLINE-2 incident), which republishes the generation
AND writes operational_status in the same step (the B-LAND-3D gotcha).
"""
import argparse
import hashlib
import io
import json
import os
import re
import sys
import zipfile
from pathlib import Path

sys.path.insert(0, "live_data")
from rmv2_live.config import load_config
from rmv2_live.pipeline import RefreshPipeline
from rmv2_live import feed_factory
from rmv2_live.adapters import normalize
from rmv2_live.offline_binding import bind_offline_current

ROOT = Path(".").resolve()
CFG_PATH = ROOT / "live_data/config/sources.v1.json"
CACHE = ROOT / "research/prefetch/forecasts"
AS_OF = "2026-08-06T10:00:00Z"
FAMILY = "philadelphia_spf"
SPF_PAGE = ("https://www.philadelphiafed.org/surveys-and-data/real-time-data-"
            "research/survey-of-professional-forecasters")
ANX_PAGE = ("https://www.philadelphiafed.org/surveys-and-data/real-time-data-"
            "research/anxious-index")

FIREWALL = (
    "FORECAST class: forecast of a MEASUREMENT -> eligible LABELED "
    "substituted_diagnostic input AND comparator; NOT an NBER recession-LABEL "
    "product (RECESS/Anxious are GDP-decline probabilities). Never a construction "
    "input unadmitted. family philadelphia_spf. B-LAND-8."
)


def _manifest_index(subdir, manifest_name="MANIFEST.sha256.json"):
    man = json.load(open(CACHE / subdir / manifest_name))
    idx = {}
    for f in man["files"]:
        idx[f["path"].split("/")[-1]] = f
    return man.get("generated_utc"), idx


def _cache_manifest(subdir, filename, gen_utc, idx, url):
    path = CACHE / subdir / filename
    body = path.read_bytes()
    sha = hashlib.sha256(body).hexdigest()
    row = idx.get(filename)
    if row is None:
        raise SystemExit("NO MANIFEST ROW for %s/%s" % (subdir, filename))
    if sha != row["sha256"] or len(body) != row["bytes"]:
        raise SystemExit("CACHE MISMATCH %s: sha/len vs manifest" % filename)
    manifest = {
        "fetch_utc": gen_utc,
        "source_bytes_length": len(body),
        "source_sha256": sha,
        "url": url,
    }
    return body, manifest


def _detect_sheet(body):
    """First worksheet name (repaired in-memory; original bytes untouched)."""
    from rmv2_live.forecast_xlsx import repair_core_xml_bytes
    import openpyxl
    wb = openpyxl.load_workbook(repair_core_xml_bytes(body), read_only=True)
    try:
        for sn in wb.sheetnames:
            ws = wb[sn]
            if hasattr(ws, "iter_rows"):
                return sn
        return wb.sheetnames[0]
    finally:
        wb.close()


def _agg_source(filename, sheet, statistic, transform):
    stem = filename[:-5]  # drop .xlsx
    unit = ("SPF %s annualized growth forecast, percent" % statistic.lower()
            if transform == "growth"
            else "SPF %s level forecast (units per SPF documentation; see varcode)"
            % statistic.lower())
    return {
        "adapter": "forecast_xlsx",
        "allowed_hosts": ["www.philadelphiafed.org"],
        "archival": True,
        "coverage_source_ids": [FAMILY],
        "enabled": False,
        "endpoint": "%s#%s" % (SPF_PAGE, filename),
        "expected_content_types": [
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
        ],
        "frequency": "quarterly",
        "information_set_mode": "substituted_diagnostic",
        "label": (
            "SPF %s %s forecasts (%s) [OFFLINE-CURRENT frozen forecast snapshot; "
            "CH-R35 cache; %s]" % (statistic, transform, filename, FIREWALL)
        ),
        "max_bytes": 33554432,
        "method_version": "spf_aggregate_forecast_offline.v1",
        "poll_seconds": 86400,
        "publisher": "Federal Reserve Bank of Philadelphia",
        "publisher_release_clock": (
            "quarterly named survey vintage (released after the advance GDP "
            "release); exact media download URL not captured by CH-R35"
        ),
        "rights_status": "published_survey_data",
        "secret_env": None,
        "secret_required": False,
        "series": {
            "forecast_shape": "spf_aggregate",
            "sheet": sheet,
            "statistic": statistic,
            "id_prefix": "SPF_",
            "unit": unit,
            "label": "SPF %s %s forecast. %s" % (statistic, transform, FIREWALL),
        },
        "source_id": "philadelphia_spf_%s" % stem,
        "value_status": "actual",
    }


def _anxious_source(filename, sheet):
    return {
        "adapter": "forecast_xlsx",
        "allowed_hosts": ["www.philadelphiafed.org"],
        "archival": True,
        "coverage_source_ids": [FAMILY],
        "enabled": False,
        "endpoint": "%s#%s" % (ANX_PAGE, filename),
        "expected_content_types": [
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
        ],
        "frequency": "quarterly",
        "information_set_mode": "substituted_diagnostic",
        "label": (
            "Philadelphia Fed Anxious Index (mean SPF probability of a negative "
            "quarter-over-quarter real GDP growth, one quarter ahead) [OFFLINE-"
            "CURRENT frozen snapshot; CH-R35 cache; %s]" % FIREWALL
        ),
        "max_bytes": 33554432,
        "method_version": "spf_anxious_index_offline.v1",
        "poll_seconds": 86400,
        "publisher": "Federal Reserve Bank of Philadelphia",
        "publisher_release_clock": (
            "quarterly named survey vintage; exact media download URL not "
            "captured by CH-R35"
        ),
        "rights_status": "published_survey_data",
        "secret_env": None,
        "secret_required": False,
        "series": {
            "forecast_shape": "anxious_index",
            "sheet": sheet,
            "series_id": "SPF_ANXIOUS_INDEX",
            "unit": "percent probability",
            "label": ("Anxious Index (mean prob of negative q/q real GDP growth, "
                      "one quarter ahead). %s" % FIREWALL),
        },
        "source_id": "philadelphia_spf_anxious_index",
        "value_status": "actual",
    }


def build_plan():
    plan = []  # (src, body, manifest)
    # Anxious
    gen_ax, idx_ax = _manifest_index("anxious_index")
    ax_file = "anxious_index_chart.xlsx"
    body, man = _cache_manifest("anxious_index", ax_file, gen_ax, idx_ax,
                                "%s#%s" % (ANX_PAGE, ax_file))
    plan.append((_anxious_source(ax_file, _detect_sheet(body)), body, man))
    # SPF aggregates: mean_/median_ files only (individuals + micro master deferred)
    gen_spf, idx_spf = _manifest_index("spf")
    agg = sorted(f for f in os.listdir(CACHE / "spf")
                 if f.startswith(("mean_", "median_")) and f.endswith(".xlsx"))
    for filename in agg:
        statistic = "MEAN" if filename.startswith("mean_") else "MEDIAN"
        body, man = _cache_manifest("spf", filename, gen_spf, idx_spf,
                                    "%s#%s" % (SPF_PAGE, filename))
        sheet = _detect_sheet(body)
        transform = "growth" if sheet.lower().endswith("growth") else "level"
        plan.append((_agg_source(filename, sheet, statistic, transform), body, man))
    return plan


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    a = ap.parse_args()

    plan = build_plan()
    cfg = load_config(CFG_PATH)
    existing = {s["source_id"] for s in cfg["sources"]}
    for src, _b, _m in plan:
        if src["source_id"] in existing:
            raise SystemExit("COLLISION: %s already in config" % src["source_id"])

    pipeline = RefreshPipeline(ROOT, cfg)
    pipeline.store.initialize()

    # clobber-proof: gate EVERY candidate against the live config BEFORE any write,
    # threading each accepted candidate into the running config so intra-batch
    # collisions are caught too. Also parse each for the plan report.
    running = dict(cfg)
    running["sources"] = list(cfg["sources"])
    report = {}
    all_series = set()
    for src, body, _m in plan:
        feed_factory._validate_source_candidate(ROOT, running, src)
        recs = normalize(src, body, AS_OF)
        avail = sum(1 for r in recs if r.get("value_status") != "unavailable")
        sids = sorted(set(r["series_id"] for r in recs))
        dup = all_series.intersection(sids)
        if dup:
            raise SystemExit("SERIES COLLISION across files: %r" % sorted(dup))
        all_series.update(sids)
        periods = [r["observation_period"] for r in recs if r["observation_period"]]
        report[src["source_id"]] = {
            "records": len(recs), "available": avail, "n_series": len(sids),
            "first": min(periods), "last": max(periods),
        }
        running["sources"] = list(running["sources"]) + [src]
    print("GATE PASS all %d sources, %d distinct series." % (len(plan), len(all_series)))
    print("total records:", sum(v["records"] for v in report.values()))

    if a.dry:
        json.dump({"report": report, "n_series": len(all_series)},
                  open("research/BLAND8_land_plan.json", "w"), indent=1, sort_keys=True)
        print("DRY: gated + parsed, no writes.")
        return

    bound = {}
    for src, body, manifest in plan:
        oc = bind_offline_current(pipeline, src, body, manifest, AS_OF)
        bound[src["source_id"]] = {
            "records": oc["record_count"], "receipt": oc["receipt_sha256"],
            "latest": oc["latest_observation_period"],
        }
        print("BOUND", src["source_id"], oc["record_count"], oc["receipt_sha256"][:12])

    raw = json.load(open(CFG_PATH))
    raw["sources"] = list(raw["sources"]) + [src for src, _b, _m in plan]
    tmp = str(CFG_PATH) + ".tmp"
    json.dump(raw, open(tmp, "w"), indent=1, sort_keys=True)
    os.replace(tmp, CFG_PATH)
    print("CONFIG written; n_sources", len(raw["sources"]),
          "archival", sum(1 for s in raw["sources"] if s.get("archival")))

    ids = [src["source_id"] for src, _b, _m in plan]
    cfg2 = load_config(CFG_PATH)
    pipe2 = RefreshPipeline(ROOT, cfg2)
    result = pipe2.refresh(source_ids=ids, due_only=False)
    outcomes = result.get("outcomes") if isinstance(result, dict) else None
    pointer = result.get("pointer") if isinstance(result, dict) else None
    print("TARGETED REFRESH outcomes:", json.dumps(outcomes))
    print("POINTER generation:", (pointer or {}).get("generation_sha256", "<none>"))
    json.dump({"bound": bound, "report": report,
               "refresh": {"outcomes": outcomes, "pointer": pointer},
               "n_series": len(all_series)},
              open("research/BLAND8_land_out.json", "w"), indent=1, default=str,
              sort_keys=True)


main()
