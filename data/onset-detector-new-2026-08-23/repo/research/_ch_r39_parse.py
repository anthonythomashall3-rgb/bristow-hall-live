import openpyxl, json, os, glob, re, time
base="research/prefetch/rtdsm"
rows=[]
t0=time.time()
for xf in sorted(glob.glob(base+"/*/*.xlsx")):
    var=xf.split("/")[-2]; fn=os.path.basename(xf)
    byt=os.path.getsize(xf)
    try:
        wb=openpyxl.load_workbook(xf, read_only=True, data_only=True)
        ws=wb[wb.sheetnames[0]]
        it=ws.iter_rows(values_only=True)
        hdr=next(it)
        vcols=[str(h) for h in hdr[1:] if h not in (None,"")]
        nv=len(vcols)
        first_v=vcols[0] if vcols else None
        last_v=vcols[-1] if vcols else None
        dfirst=dlast=None; recs=0
        for r in it:
            d=r[0]
            if d in (None,""): continue
            if dfirst is None: dfirst=str(d)
            dlast=str(d)
            for c in r[1:1+nv]:
                if c is None: continue
                s=str(c).strip()
                if s and s!="#N/A" and s.upper()!="NA": recs+=1
        wb.close()
        rows.append(dict(var=var,file=fn,bytes=byt,vintages=nv,
            first_vintage=first_v,last_vintage=last_v,
            obs_first=dfirst,obs_last=dlast,records=recs))
    except Exception as e:
        rows.append(dict(var=var,file=fn,bytes=byt,error=str(e)[:80]))
json.dump(rows, open("research/_ch_r39_parsed.json","w"), indent=0)
print("files:",len(rows),"elapsed:",round(time.time()-t0,1),"s")
print("errors:",sum(1 for r in rows if "error" in r))
