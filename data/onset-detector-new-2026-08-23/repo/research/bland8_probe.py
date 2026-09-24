import io, re, zipfile, openpyxl
def repair(b):
    zin=zipfile.ZipFile(io.BytesIO(b)); items={n:zin.read(n) for n in zin.namelist()}
    cx=items.get('docProps/core.xml')
    if cx: items['docProps/core.xml']=re.sub(rb'T\s?(\d):', rb'T0\1:', cx)
    out=io.BytesIO(); zo=zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED)
    for n,d in items.items(): zo.writestr(n,d)
    zo.close(); out.seek(0); return out
def dump(path,label,data_only=True):
    b=open(path,'rb').read()
    wb=openpyxl.load_workbook(repair(b),read_only=True,data_only=data_only)
    print("====",label,path.split('/')[-1],"data_only=",data_only,"sheets:",wb.sheetnames)
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    for sn in wb.sheetnames:
        ws=wb[sn]
        if not isinstance(ws,ReadOnlyWorksheet): 
            print(" [chartsheet]",repr(sn)); continue
        rows=list(ws.iter_rows(values_only=True))
        nonnull=sum(1 for r in rows for c in r if c is not None)
        print(" sheet",repr(sn),"nrows",len(rows),"nonnull_cells",nonnull)
        for r in rows[:6]: print("   ",r[:14])
        print("   last:",rows[-1][:14] if rows else None)
    wb.close()
dump("research/prefetch/forecasts/anxious_index/anxious_index_chart.xlsx","ANXIOUS-CHART")
dump("research/prefetch/forecasts/anxious_index/anxious_index_chart.xlsx","ANXIOUS-CHART-FORMULA",False)
dump("research/prefetch/forecasts/spf/mean_rgdp_growth.xlsx","AGG-MEAN")
